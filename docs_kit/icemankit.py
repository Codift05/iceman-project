# -*- coding: utf-8 -*-
"""Icemankit: pembangun dokumen proyek Iceman.

Menghasilkan dokumen .docx dan .pdf yang seragam: A4, Arial, teks justify,
monokrom, cover bergaya Iceman, daftar isi dan daftar gambar bernomor halaman
terverifikasi, serta footer berisi nomor halaman saja.

Pemakaian singkat:

    from icemankit import Doc
    d = Doc(code="SRS", title="Software Requirements Specification")
    d.h1("1. Pendahuluan")
    d.p("Isi paragraf.")
    d.table([["Kolom A", "Kolom B"], ["1", "2"]])
    d.build("02_SRS_Iceman_Apps")
"""
import os, re, shutil, subprocess, sys
from docx import Document
from docx.shared import Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

import cover as coverlib

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
BUILD = os.path.join(HERE, "build")
TEMPLATE = os.path.join(ASSETS, "template.docx")
PLANTUML = os.path.join(ASSETS, "plantuml.jar")
THEME = os.path.join(ASSETS, "theme.iuml")
OUTDIR = os.path.abspath(os.path.join(HERE, ".."))

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TEXT_TW = 9746                                   # lebar blok teks A4 (twips)
USABLE_EMU = int(TEXT_TW / 1440 * 914400)
MAX_IMG_H_EMU = int(9.6 * 914400)
CELL_SZ = "18"                                   # 9pt
HDR_FILL = "D9D9D9"
LBL_FILL = "F2F2F2"

STYLE_ID = {"Heading 1": "Heading1", "Heading 2": "Heading2", "Heading 3": "Heading3",
            "Normal": "Normal", "List Bullet": "ListBullet", "Caption": "Caption",
            "Requirement ID": "RequirementID", "Small Note": "SmallNote"}


# --------------------------------------------------------------- util XML
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def _run(text, bold=False, italic=False, sz=None, mono=False):
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    fam = "Courier New" if mono else "Arial"
    rPr.append(_el("w:rFonts", ascii=fam, hAnsi=fam, cs=fam, eastAsia=fam))
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    rPr.append(_el("w:color", val="000000"))
    if sz:
        rPr.append(_el("w:sz", val=sz))
    r.append(rPr)
    if text == "\t":
        r.append(OxmlElement("w:tab"))
    else:
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
    return r


def _para(text="", style=None, align=None, bold=False, italic=False, sz=None, keep_next=False):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    if style:
        sid = STYLE_ID.get(style)
        if sid is None:
            raise KeyError(f"style tidak dikenal: {style!r}")
        pPr.append(_el("w:pStyle", val=sid))
    if keep_next:
        pPr.append(OxmlElement("w:keepNext"))
    if align:
        pPr.append(_el("w:jc", val=align))
    if len(pPr):
        p.append(pPr)
    if text:
        p.append(_run(text, bold=bold, italic=italic, sz=sz))
    return p


def _tc(lines, width, fill=None, bold=False, valign="top"):
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tcPr.append(_el("w:tcW", type="dxa", w=str(width)))
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", "70"), ("start", "90"), ("bottom", "70"), ("end", "90")):
        mar.append(_el("w:" + side, w=val, type="dxa"))
    tcPr.append(mar)
    tcPr.append(_el("w:vAlign", val=valign))
    tc.append(tcPr)
    if isinstance(lines, str):
        lines = [lines]
    for ln in lines:
        p = OxmlElement("w:p")
        p.append(_run(str(ln), bold=bold, sz=CELL_SZ))
        tc.append(p)
    return tc


def _table(rows, widths=None, header=True, label_col=False):
    ncol = max(len(r) for r in rows)
    if widths is None:
        widths = [TEXT_TW // ncol] * ncol
    widths = list(widths)
    widths[-1] += TEXT_TW - sum(widths)

    tbl = OxmlElement("w:tbl")
    tblPr = OxmlElement("w:tblPr")
    tblPr.append(_el("w:tblStyle", val="TableGrid"))
    tblPr.append(_el("w:tblW", type="dxa", w=str(TEXT_TW)))
    tblPr.append(_el("w:tblInd", w="0", type="dxa"))
    tblPr.append(_el("w:tblLayout", type="fixed"))
    look = OxmlElement("w:tblLook")
    for k, v in (("firstColumn", "1"), ("firstRow", "1"), ("lastColumn", "0"),
                 ("lastRow", "0"), ("noHBand", "0"), ("noVBand", "1"), ("val", "04A0")):
        look.set(qn("w:" + k), v)
    tblPr.append(look)
    tbl.append(tblPr)

    grid = OxmlElement("w:tblGrid")
    for w in widths:
        grid.append(_el("w:gridCol", w=str(w)))
    tbl.append(grid)

    for i, row in enumerate(rows):
        tr = OxmlElement("w:tr")
        is_hdr = header and i == 0
        if is_hdr:
            trPr = OxmlElement("w:trPr")
            trPr.append(OxmlElement("w:tblHeader"))
            trPr.append(OxmlElement("w:cantSplit"))
            tr.append(trPr)
        for j in range(ncol):
            cell = row[j] if j < len(row) else ""
            if is_hdr:
                tr.append(_tc(cell, widths[j], fill=HDR_FILL, bold=True, valign="center"))
            elif label_col and j == 0:
                tr.append(_tc(cell, widths[j], fill=LBL_FILL, bold=True))
            else:
                tr.append(_tc(cell, widths[j]))
        tbl.append(tr)
    return tbl


# --------------------------------------------------------------- PlantUML
def render_puml(puml_path, out_dir=None):
    """Render satu berkas .puml menjadi PNG memakai tema Iceman."""
    out_dir = out_dir or os.path.join(BUILD, "img")
    os.makedirs(out_dir, exist_ok=True)
    src_dir = os.path.dirname(os.path.abspath(puml_path))
    theme_local = os.path.join(src_dir, "_theme.iuml")
    if not os.path.exists(theme_local):
        shutil.copy(THEME, theme_local)
    subprocess.run(["java", "-jar", PLANTUML, "-tpng", "-Sdpi=160",
                    "-o", os.path.abspath(out_dir), os.path.abspath(puml_path)],
                   check=True, capture_output=True)
    name = os.path.splitext(os.path.basename(puml_path))[0] + ".png"
    return os.path.join(out_dir, name)


def image_fit(path):
    """Lebar dan tinggi EMU agar gambar pas pada blok teks tanpa melebihi halaman."""
    with Image.open(path) as im:
        w, h = im.size
    nw = USABLE_EMU
    nh = int(nw * h / w)
    if nh > MAX_IMG_H_EMU:
        nh = MAX_IMG_H_EMU
        nw = int(nh * w / h)
    return nw, nh, w / h


# --------------------------------------------------------------- dokumen
class Doc:
    def __init__(self, code, title, product="ICEMAN APPS",
                 subtitle="Aplikasi Pemesanan dan Distribusi Es Kristal",
                 version="1.0", status="Draft untuk Validasi",
                 basis="Disusun berdasarkan PRD Iceman Apps v1.1"):
        self.code, self.title, self.product = code, title, product
        self.subtitle, self.version, self.status, self.basis = subtitle, version, status, basis
        self.blocks = []          # elemen XML siap sisip
        self.figures = []         # (caption, path)
        self._fig_no = 0

    # ---------- blok teks ----------
    def h1(self, text):
        self.blocks.append(_para(text, style="Heading 1")); return self

    def h2(self, text):
        self.blocks.append(_para(text, style="Heading 2")); return self

    def h3(self, text):
        self.blocks.append(_para(text, style="Heading 3")); return self

    def p(self, text):
        self.blocks.append(_para(text, style="Normal", align="both")); return self

    def bullets(self, items):
        for it in items:
            self.blocks.append(_para(it, style="List Bullet", align="both"))
        return self

    def note(self, text):
        self.blocks.append(_para(text, style="Small Note", align="both")); return self

    def req(self, ident, pairs):
        """Blok requirement: judul ber-ID + tabel dua kolom berlabel."""
        self.blocks.append(_para(ident, style="Requirement ID"))
        self.blocks.append(_table(pairs, [3100, 6646], header=False, label_col=True))
        self.blocks.append(_para(""))
        return self

    def table(self, rows, widths=None, header=True, label_col=False):
        self.blocks.append(_table(rows, widths, header, label_col))
        self.blocks.append(_para(""))
        return self

    def codeblock(self, lines, caption=None):
        """Blok kode monospace berlatar abu, untuk SQL atau potongan konfigurasi."""
        if isinstance(lines, str):
            lines = lines.split("\n")
        tc_lines = [ln.replace("\t", "    ") for ln in lines]
        tbl = OxmlElement("w:tbl")
        tblPr = OxmlElement("w:tblPr")
        tblPr.append(_el("w:tblStyle", val="TableGrid"))
        tblPr.append(_el("w:tblW", type="dxa", w=str(TEXT_TW)))
        tblPr.append(_el("w:tblInd", w="0", type="dxa"))
        tblPr.append(_el("w:tblLayout", type="fixed"))
        tbl.append(tblPr)
        grid = OxmlElement("w:tblGrid")
        grid.append(_el("w:gridCol", w=str(TEXT_TW)))
        tbl.append(grid)
        tr = OxmlElement("w:tr")
        tc = OxmlElement("w:tc")
        tcPr = OxmlElement("w:tcPr")
        tcPr.append(_el("w:tcW", type="dxa", w=str(TEXT_TW)))
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F7F7F7")
        tcPr.append(shd)
        mar = OxmlElement("w:tcMar")
        for side, val in (("top", "90"), ("start", "120"), ("bottom", "90"), ("end", "120")):
            mar.append(_el("w:" + side, w=val, type="dxa"))
        tcPr.append(mar)
        tc.append(tcPr)
        for ln in tc_lines:
            p = OxmlElement("w:p")
            pPr = OxmlElement("w:pPr")
            pPr.append(_el("w:spacing", before="0", after="0", line="240", lineRule="auto"))
            p.append(pPr)
            if ln.strip():
                p.append(_run(ln, sz="16", mono=True))
            tc.append(p)
        tr.append(tc); tbl.append(tr)
        self.blocks.append(tbl)
        if caption:
            self.blocks.append(_para(caption, style="Small Note", align="both"))
        self.blocks.append(_para(""))
        return self

    def spacer(self):
        self.blocks.append(_para("")); return self

    def pagebreak(self):
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        r.append(_el("w:br", type="page"))
        p.append(r)
        self.blocks.append(p); return self

    # ---------- gambar ----------
    def figure(self, img_path, caption):
        """Sisipkan gambar; penomoran 'Gambar N.' ditambahkan otomatis."""
        self._fig_no += 1
        cap = f"Gambar {self._fig_no}. {caption}"
        self.blocks.append(("IMG", img_path, cap))
        self.figures.append(cap)
        return self

    def diagram(self, puml_path, caption):
        return self.figure(render_puml(puml_path), caption)

    # ---------- perakitan ----------
    def _cover_png(self, slug):
        meta = [f"Versi {self.version}  |  {self.status}"]
        if self.basis:
            meta.append(self.basis)
        return coverlib.make_cover(
            code=self.code, title=self.title, product=self.product,
            subtitle=self.subtitle, meta_lines=meta,
            out_path=os.path.join(BUILD, f"cover_{slug}.png"))

    def _assemble(self, docx_path, cover_png):
        shutil.copy(TEMPLATE, docx_path)
        doc = Document(docx_path)
        body = doc.element.body
        holder = body.find(W + "p")          # paragraf pembawa sectPr cover

        # gambar cover, penuh satu halaman
        tmp = doc.add_paragraph(); tmp.alignment = 1
        with Image.open(cover_png) as im:
            cw, ch = im.size
        pw = int(11906 / 1440 * 914400)
        phh = int(16838 / 1440 * 914400)
        ph = int(pw * ch / cw)
        if ph > phh:
            ph = phh; pw = int(ph * cw / ch)
        tmp.add_run().add_picture(cover_png, width=Emu(pw), height=Emu(ph))
        cp = tmp._element
        cp.getparent().remove(cp)
        cpPr = cp.find(W + "pPr")
        if cpPr is None:
            cpPr = OxmlElement("w:pPr")
            cp.insert(0, cpPr)
        cpPr.insert(0, _el("w:spacing", before="0", after="0", line="240", lineRule="auto"))
        holder.addprevious(cp)

        # isi dokumen
        anchor = holder
        for blk in self.blocks:
            if isinstance(blk, tuple) and blk[0] == "IMG":
                _, path, cap = blk
                nw, nh, _ = image_fit(path)
                t2 = doc.add_paragraph(); t2.alignment = 1
                t2.add_run().add_picture(path, width=Emu(nw), height=Emu(nh))
                pel = t2._element
                pel.getparent().remove(pel)
                pPr = pel.find(W + "pPr")
                if pPr is None:
                    pPr = OxmlElement("w:pPr"); pel.insert(0, pPr)
                pPr.insert(0, OxmlElement("w:keepNext"))
                anchor.addnext(pel); anchor = pel
                capel = _para(cap, style="Caption", align="center")
                anchor.addnext(capel); anchor = capel
                sp = _para("")
                anchor.addnext(sp); anchor = sp
            else:
                anchor.addnext(blk); anchor = blk
        doc.save(docx_path)

    # ---------- daftar isi ----------
    def _toc_lines(self, doc):
        entries, figs = [], []
        for p in doc.paragraphs:
            st, txt = p.style.name, p.text.strip()
            if not txt:
                continue
            if st == "Heading 1":
                entries.append((txt, 1))
            elif st == "Heading 2":
                entries.append((txt, 2))
            elif st == "Caption" and txt.startswith("Gambar"):
                figs.append(txt)
        return entries, figs

    @staticmethod
    def _clip_entry(txt, limit=92):
        """Pendekkan entri daftar agar tidak membungkus ke baris kedua."""
        txt = txt.rstrip(".")
        if len(txt) <= limit:
            return txt
        cut = txt[:limit].rsplit(" ", 1)[0]
        return cut.rstrip(" ,;:") + "..."

    @staticmethod
    def _toc_line(title, level):
        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tabs.append(_el("w:tab", val="right", leader="dot", pos=str(TEXT_TW)))
        pPr.append(tabs)
        if level == 2:
            pPr.append(_el("w:ind", left="340"))
        pPr.append(_el("w:spacing", before="0", after="40", line="240", lineRule="auto"))
        p.append(pPr)
        p.append(_run(title, bold=(level == 1), sz="20"))
        p.append(_run("\t", sz="20"))
        p.append(_run("0", bold=(level == 1), sz="20"))
        return p

    def _insert_toc(self, docx_path):
        doc = Document(docx_path)
        entries, figs = self._toc_lines(doc)
        first_h1 = next((p._element for p in doc.paragraphs if p.style.name == "Heading 1"), None)
        if first_h1 is None:
            doc.save(docx_path); return
        blocks = [_para("Daftar Isi", style="Heading 1")]
        blocks += [self._toc_line(self._clip_entry(t), l) for t, l in entries]
        if figs:
            blocks.append(self._pb())
            blocks.append(_para("Daftar Gambar", style="Heading 1"))
            blocks += [self._toc_line(self._clip_entry(f), 2) for f in figs]
        blocks.append(self._pb())
        for b in blocks:
            first_h1.addprevious(b)
        doc.save(docx_path)

    @staticmethod
    def _pb():
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        r.append(_el("w:br", type="page"))
        p.append(r)
        return p

    # ---------- PDF dan nomor halaman ----------
    @staticmethod
    def _to_pdf(docx_path, out_dir):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                        "--outdir", out_dir, docx_path],
                       check=True, capture_output=True, timeout=600)
        return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")

    @staticmethod
    def _pdf_pages(pdf_path):
        txt = pdf_path + ".txt"
        subprocess.run(["pdftotext", "-layout", pdf_path, txt], check=True, capture_output=True)
        raw = open(txt, encoding="utf-8").read().split("\f")
        out = []
        for pg in raw:
            lines = []
            for ln in pg.split("\n"):
                if "...." in ln:
                    continue
                ln = re.sub(r"\s+", " ", ln).strip().lower()
                if ln:
                    lines.append(ln)
            out.append(lines)
        return out

    def _fill_pages(self, docx_path, pdf_path, offset=1):
        pages = self._pdf_pages(pdf_path)
        doc = Document(docx_path)
        rows = []
        for p in doc.paragraphs:
            pPr = p._element.find(W + "pPr")
            if pPr is None:
                continue
            tabs = pPr.find(W + "tabs")
            if tabs is not None and any(t.get(qn("w:leader")) == "dot" for t in tabs):
                rows.append(p)
        cursor, filled, miss = 0, 0, 0
        for p in rows:
            title = re.sub(r"\s+", " ", p.runs[0].text).strip().lower()
            probe = title.split(".")[0] + "." if title.startswith("gambar") else title[:38]
            found = None
            for pi in range(cursor, len(pages)):
                if any(ln.startswith(probe) for ln in pages[pi]):
                    found = pi; break
            if found is None:
                for pi in range(1, len(pages)):
                    if any(ln.startswith(probe) for ln in pages[pi]):
                        found = pi; break
            if found is None:
                miss += 1; continue
            cursor = found
            p.runs[-1].text = str(found + 1 - offset)
            filled += 1
        doc.save(docx_path)
        return filled, miss

    # ---------- entri utama ----------
    def build(self, slug, verbose=True):
        os.makedirs(BUILD, exist_ok=True)
        docx_path = os.path.join(OUTDIR, slug + ".docx")
        cover_png = self._cover_png(slug)
        self._assemble(docx_path, cover_png)
        self._insert_toc(docx_path)

        pdf = self._to_pdf(docx_path, BUILD)                 # jalan pertama: ukur halaman
        filled, miss = self._fill_pages(docx_path, pdf)
        pdf = self._to_pdf(docx_path, BUILD)                 # jalan kedua: nomor final
        shutil.copy(pdf, os.path.join(OUTDIR, slug + ".pdf"))

        if verbose:
            n = len(self._pdf_pages(pdf)) - 1
            print(f"  {slug}: {n} halaman | daftar isi terisi {filled}"
                  f"{', gagal ' + str(miss) if miss else ''} | gambar {len(self.figures)}")
        return os.path.join(OUTDIR, slug + ".pdf")
